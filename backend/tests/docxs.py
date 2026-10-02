"""Builds small but real Word files for tests."""

import io

import docx


def make_docx(*blocks) -> bytes:
    """Each block is one of:
    ("heading", "Payment terms")         a Heading 1 paragraph
    ("p", "Payment is due monthly.")     a normal paragraph
    ("table", [["Item", "Price"], ...])  a table, one list per row
    """
    document = docx.Document()
    for kind, content in blocks:
        if kind == "heading":
            document.add_heading(content, level=1)
        elif kind == "p":
            document.add_paragraph(content)
        elif kind == "table":
            table = document.add_table(rows=len(content), cols=len(content[0]))
            for row, values in zip(table.rows, content):
                for cell, value in zip(row.cells, values):
                    cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
